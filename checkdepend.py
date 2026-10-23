#!/usr/bin/env python3
"""
checkdepend.py

With no arguments, ask for a Linux upstream tree and commit IDs, switch that
tree to master, pull the latest code, and list commits that fix each one.

--backport keeps the original dependency check between a local kernel tree
and a Torvalds tree.

--install, --uninstall, and --reinstall compile this program into a binary
and install or remove it.

After either check finishes, the tool can write the results to an Excel sheet.
"""
import atexit
import argparse
import datetime
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

# Marker stored in the compiled binary so uninstall only removes this tool.
TOOL_MARKER = "CHECKDEPEND_TOOL_V1"
INSTALL_NAME = "checkdepend"
STAMP = Path.home() / ".config" / "checkdepend" / "install-path"

FIXES_RE = re.compile(
    r"^\s*Fixes:\s*(?:commit\s+)?([0-9a-fA-F]{7,40})\b",
    re.IGNORECASE,
)

# ---------- utils ----------
def run_cmd(cmd: str, cwd: str = None) -> str:
    """Run shell command and return stdout as latin-1 string (safe for kernel logs)."""
    result = subprocess.run(cmd, cwd=cwd, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, shell=True)
    return result.stdout.decode("latin-1", errors="replace")

def write_file(path: str, content: str):
    with open(path, "w", encoding="latin-1", errors="replace") as f:
        f.write(content)

def append_file(path: str, content: str):
    with open(path, "a", encoding="latin-1", errors="replace") as f:
        f.write(content)

def file_read(path: str) -> str:
    with open(path, "r", encoding="latin-1", errors="replace") as f:
        return f.read()

def program_name() -> str:
    base = os.path.basename(__file__)
    if base.startswith("checkdepend-embed-") or not base.endswith(".py"):
        return INSTALL_NAME
    return "checkdepend.py"

def cleanup_embedded_script():
    """Remove the temporary copy created by the compiled binary."""
    name = os.path.basename(__file__)
    if name.startswith("checkdepend-embed-"):
        try:
            os.unlink(__file__)
        except OSError:
            pass

atexit.register(cleanup_embedded_script)

def prompt_line(text: str = ""):
    """Read one line. Return None on end of input."""
    try:
        if text:
            return input(text)
        return input()
    except EOFError:
        print()
        return None

def git_run(repo: str, *args: str):
    """Run git in repo. Return (code, stdout, stderr) decoded as latin-1."""
    result = subprocess.run(
        ["git", "-C", repo, *args],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    out = result.stdout.decode("latin-1", errors="replace")
    err = result.stderr.decode("latin-1", errors="replace")
    return result.returncode, out, err

# ---------- git helpers ----------
def git_show_full_commit(user_repo: str, commitish: str):
    """
    Return (full_hash, subject) for a commit-ish in user_repo.
    If lookup fails, return (None, None).
    """
    out = run_cmd(f'git show --pretty=format:"%H%n%s" --no-patch {commitish}', cwd=user_repo)
    if not out:
        return (None, None)
    parts = out.splitlines()
    if len(parts) >= 2:
        return (parts[0].strip(), parts[1].strip())
    elif len(parts) == 1:
        return (parts[0].strip(), "")
    return (None, None)

def git_find_stable_commits_mentioning(short7: str, stable_repo: str):
    """
    Returns a list of tuples for stable commits that mention short7:
      [(stable_full_hash, subject, body), ...]
    Uses `git log --all --grep=short7 -i` to get candidate commits, then parses them.
    """
    # Use an unambiguous separator so we can split results safely.
    fmt = "%H%x01%s%x01%b%x02"
    cmd = f"git log --all --grep={short7} -i --pretty=format:'{fmt}'"
    out = run_cmd(cmd, cwd=stable_repo)
    if not out:
        return []

    entries = out.split("\x02")
    results = []
    for e in entries:
        if not e.strip():
            continue
        parts = e.split("\x01")
        if len(parts) < 3:
            continue
        full_hash = parts[0].strip()
        subject = parts[1].strip()
        body = parts[2]
        results.append((full_hash, subject, body))
    return results

# ---------- matching helpers ----------
def short7_in_text_non_commented(short7: str, text: str) -> bool:
    """
    Return True if short7 occurs in text on a line where it is not preceded by '#'
    (i.e., not commented out).
    """
    for line in text.splitlines():
        if short7 in line:
            pos = line.find(short7)
            # if there's a '#' BEFORE pos on the same line, treat as commented-out
            if "#" in line[:pos]:
                continue
            return True
    return False

def read_commit_ids(intro: str):
    """Read commit IDs, one per line, until 'done' or end of input."""
    print(intro)
    commits = []
    while True:
        ln = prompt_line()
        if ln is None:
            break
        ln = ln.strip()
        if ln.lower() == "done":
            break
        if ln:
            commits.append(ln.split()[0].strip(",;"))
    return commits

def ensure_git_worktree(path: str):
    """Return an absolute repo path, or None after printing an error."""
    repo = os.path.abspath(os.path.expanduser(path))
    if not os.path.isdir(repo):
        print(f"Not a directory: {repo}")
        return None
    code, out, err = git_run(repo, "rev-parse", "--is-inside-work-tree")
    if code != 0 or out.strip() != "true":
        detail = err.strip() or "not a git repository"
        print(f"Not a git work tree: {repo} ({detail})")
        return None
    return repo

def update_master(repo: str) -> bool:
    """Switch to master and pull. Return False only when master cannot be checked out."""
    print("\n\033[1mSwitching to master and pulling latest code\033[0m\033[5m...\033[0m")
    code, out, err = git_run(repo, "switch", "master")
    if code != 0:
        code, out, err = git_run(repo, "checkout", "master")
    if code != 0:
        detail = (err or out).strip()
        print(f"Failed to switch to master: {detail}")
        return False
    print("  switched to master")

    code, out, err = git_run(repo, "pull", "--ff-only")
    pulled = (out + err).strip()
    if code != 0:
        if pulled:
            print(pulled)
        print("Warning: git pull failed. Searching master as it is.")
        return True
    if pulled:
        for line in pulled.splitlines():
            print(f"  {line}")
    else:
        print("  pulled latest")
    return True

def resolve_upstream(repo: str, commitish: str):
    """Return (full_hash, subject) from the upstream tree, or (None, None)."""
    code, out, err = git_run(
        repo, "show", "--no-patch", "--pretty=format:%H%n%s", commitish, "--")
    if code != 0 or not out.strip():
        return None, None
    lines = out.splitlines()
    full_hash = lines[0].strip()
    if len(full_hash) != 40:
        return None, None
    subject = lines[1].strip() if len(lines) > 1 else ""
    return full_hash, subject

def abbrev_is_commit(repo: str, abbrev: str, full_hash: str, cache: dict) -> bool:
    """True when abbrev resolves uniquely to full_hash."""
    key = abbrev.lower()
    if key not in cache:
        code, out, _err = git_run(
            repo, "rev-parse", "--verify", "--quiet", key + "^{commit}")
        cache[key] = code == 0 and out.strip().lower() == full_hash.lower()
    return cache[key]

def message_fixes_commit(repo: str, full_hash: str, message: str, cache: dict) -> bool:
    """True when message has a Fixes: trailer that names full_hash."""
    for line in message.splitlines():
        match = FIXES_RE.match(line)
        if not match:
            continue
        abbrev = match.group(1).lower()
        if "#" in line[:match.start(1)]:
            continue
        if not full_hash.lower().startswith(abbrev):
            continue
        if abbrev_is_commit(repo, abbrev, full_hash, cache):
            return True
    return False

def find_fixes(repo: str, full_hash: str):
    """Return [(full_hash, subject), ...] for commits on the current branch that fix full_hash."""
    fmt = "%H%x01%s%x01%b%x02"
    code, out, err = git_run(
        repo, "log", "--all-match", "-i",
        "--grep=Fixes:",
        "--grep=" + full_hash[:12],
        "--pretty=format:" + fmt,
    )
    if code != 0:
        detail = err.strip() or "git log failed"
        print(f"  warning: {detail}")
        return []

    found = []
    seen = set()
    cache = {}
    for entry in out.split("\x02"):
        if not entry.strip():
            continue
        parts = entry.split("\x01")
        if len(parts) < 3:
            continue
        fix_hash = parts[0].strip()
        subject = parts[1].strip()
        body = parts[2]
        if fix_hash.lower() == full_hash.lower() or fix_hash in seen:
            continue
        message = subject + "\n" + body
        if not message_fixes_commit(repo, full_hash, message, cache):
            continue
        seen.add(fix_hash)
        found.append((fix_hash, subject))
    return found

def run_fixes():
    """
    Ask for an upstream tree and commit IDs, update master, and list Fixes patches.
    Return (status, rows). rows is None when the check did not finish.
    """
    entered = prompt_line("\033[1mEnter path to Linux upstream source:\033[0m ")
    if entered is None:
        return 2, None
    repo = ensure_git_worktree(entered.strip())
    if repo is None:
        return 2, None

    commits = read_commit_ids(
        "\033[1mPaste upstream commit IDs (one per line). Type 'done' when finished:\033[0m")
    if not commits:
        print("No commits to check.")
        return 0, None

    if not update_master(repo):
        return 2, None

    print("\n\033[1mChecking Fixes patches\033[0m\033[5m...\033[0m\n")
    rows = []
    for commitish in commits:
        full_hash, subject = resolve_upstream(repo, commitish)
        if not full_hash:
            print(f"Warning: couldn't resolve commit '{commitish}' in upstream. Skipping.")
            continue
        short14 = full_hash[:14]
        print(f"\033[34mCommit:\033[0m {short14} - {subject}")
        fixes = find_fixes(repo, full_hash)
        if not fixes:
            print("  No Fixes patches found\n")
            rows.append({
                "commit": full_hash,
                "subject": subject,
                "fixes": "",
                "fixes_subject": "",
            })
            continue
        for fix_hash, fix_subject in fixes:
            print(f"    \033[35m*\033[0m \033[1;33m{fix_hash[:14]} {fix_subject}\033[0m")
            rows.append({
                "commit": full_hash,
                "subject": subject,
                "fixes": fix_hash,
                "fixes_subject": fix_subject,
            })
        print()

    print("\033[1;32mDone.\033[0m\n")
    return 0, rows

# ---------- backport (original check) ----------
def run_backport():
    """Original dependency check. Returns one excel row per result line."""
    user_repo = prompt_line("\033[1mEnter path to your kernel source:\033[0m ")
    if user_repo is None:
        return []
    user_repo = user_repo.strip()
    stable_repo = prompt_line("\033[1mEnter path to Torvalds kernel source:\033[0m ")
    if stable_repo is None:
        return []
    stable_repo = stable_repo.strip()

    print("\033[1mEnter commits (one per line). Type 'done' when finished:\033[0m")
    commits = []
    while True:
        ln = prompt_line()
        if ln is None:
            break
        ln = ln.strip()
        if ln.lower() == "done":
            break
        if ln:
            commits.append(ln)

    # output files
    stable_log_path = ".stable_log"
    user_log_path = ".user_log"
    full_commits_path = ".full_commits"
    dep_log_path = ".dep_log"

    print("\n\033[1mGenerating logs\033[0m\033[5m...\033[0m")

    # 1) save stable repo full git log (for reference)
    stable_log = run_cmd("git log", cwd=stable_repo)
    write_file(stable_log_path, stable_log)

    # 2) save user repo one-line log
    user_log = run_cmd("git log --oneline", cwd=user_repo)
    write_file(user_log_path, user_log)
    user_log_txt = user_log  # Keep in memory for checking

    # 3) resolve each provided commit in user repo to full hash+subject
    write_file(full_commits_path, "")
    resolved = []
    for c in commits:
        full_hash, subject = git_show_full_commit(user_repo, c)
        if not full_hash:
            print(f"Warning: couldn't resolve commit '{c}' in user repo. Skipping.")
            continue
        append_file(full_commits_path, f"{full_hash} {subject}\n")
        resolved.append((full_hash, subject))

    # clear dep log
    write_file(dep_log_path, "")

    print("\n\033[1mChecking dependencies\033[0m\033[5m...\033[0m\n")

    report = []
    for full_hash, subject in resolved:
        short7 = full_hash[:7]
        short14 = full_hash[:14]
        print(f"\033[34mCommit:\033[0m {short14} - {subject}")

        # Query stable repo for commits that mention short7
        candidates = git_find_stable_commits_mentioning(short7, stable_repo)

        deps_found = []
        for st_full, st_subject, st_body in candidates:
            # ignore exact self-match (same commit)
            if st_full.lower() == full_hash.lower():
                continue

            # Check that the occurrence is not commented out:
            combined_text = st_subject + "\n" + st_body
            if not short7_in_text_non_commented(short7, combined_text):
                continue

            deps_found.append((st_full, st_subject))

        # deduplicate by full hash
        unique = []
        seen = set()
        for h, s in deps_found:
            if h not in seen:
                seen.add(h)
                unique.append((h, s))

        if not unique:
            print("  \033[32mPASS\033[0m -> No dependencies found\n")
            report.append({
                "commit": full_hash,
                "subject": subject,
                "result": "PASS",
                "dependency": "",
                "dependency_subject": "",
                "state": "",
            })
            continue

        # Check each dependency and track if all are fixed
        all_fixed = True
        dep_status = []
        for st_full, st_subject in unique:
            dep14 = st_full[:14]
            dep_entry = f"{dep14} {st_subject}"
            append_file(dep_log_path, dep_entry + "\n")

            # Check if this dependency is already fixed in user repo
            is_fixed = bool(re.search(re.escape(st_subject), user_log_txt))
            dep_status.append((st_full, st_subject, dep_entry, is_fixed))
            if not is_fixed:
                all_fixed = False

        # Print result based on whether all dependencies are fixed
        if all_fixed:
            print("  \033[32mPASS\033[0m -> No dependencies found\n")
            result = "PASS"
        else:
            print("  \033[31mFAIL\033[0m -> new bugfix needed")
            for _st_full, _st_subject, dep_entry, is_fixed in dep_status:
                if is_fixed:
                    print(f"    \033[35m*\033[0m \033[1;33m{dep_entry}\033[0m \033[32m-> Fixed\033[0m")
                else:
                    print(f"    \033[35m*\033[0m \033[1,33m{dep_entry}\033[0m")
            print()
            result = "FAIL"

        for st_full, st_subject, _dep_entry, is_fixed in dep_status:
            report.append({
                "commit": full_hash,
                "subject": subject,
                "result": result,
                "dependency": st_full,
                "dependency_subject": st_subject,
                "state": "Fixed" if is_fixed else "needed",
            })

    print("\033[1,32mDone.\033[0m\n")
    return report

# ---------- excel ----------
def excel_text(value) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text

def write_excel(rows, mode: str) -> str:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if mode == "fixes":
        headers = ["Commit", "Subject", "Fixes commit", "Fixes subject"]
        keys = ["commit", "subject", "fixes", "fixes_subject"]
        title = "Fixes"
        widths = [42, 78, 42, 78]
    else:
        headers = ["Commit", "Subject", "Result", "Dependency",
                   "Dependency subject", "State"]
        keys = ["commit", "subject", "result", "dependency",
                "dependency_subject", "state"]
        title = "Backport"
        widths = [42, 78, 14, 42, 78, 16]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    body_font = Font(name="Calibri", size=11)
    align = Alignment(vertical="center")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align
    for row_index, row in enumerate(rows, 2):
        for col, key in enumerate(keys, 1):
            cell = ws.cell(row_index, col, excel_text(row.get(key, "")))
            cell.font = body_font
            cell.alignment = align
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    last = max(1, len(rows) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    path = os.path.abspath(f"checkdepend-{mode}-{stamp}.xlsx")
    wb.save(path)
    return path

def ask_excel(rows, mode: str):
    """Ask once, after checking. Enter or anything other than yes skips the sheet."""
    if not rows:
        return
    answer = prompt_line("\033[1mGenerate excel sheet? [y/N]:\033[0m ")
    if answer is None:
        return
    if answer.strip().lower() not in ("y", "yes"):
        print("No excel sheet written.")
        return
    try:
        path = write_excel(rows, mode)
    except ImportError:
        print("openpyxl is not installed, so no excel sheet was written.")
        return
    except OSError as exc:
        print(f"Could not write excel sheet: {exc}")
        return
    print(f"Wrote {path}")

# ---------- compile and install ----------
def c_string(text: str) -> str:
    escaped = text.replace("\\", "\\\\").replace('"', '\\"')
    return f'"{escaped}"'

def render_c(source: bytes, interpreter: str) -> str:
    lines = [
        "/* generated by checkdepend --install */",
        "#define _GNU_SOURCE",
        "#include <stdio.h>",
        "#include <stdlib.h>",
        "#include <unistd.h>",
        "",
        "__attribute__((used))",
        f"static const char tool_marker[] = {c_string(TOOL_MARKER)};",
        f"static const char interpreter[] = {c_string(interpreter)};",
        "static const unsigned char script[] = {",
    ]
    for offset in range(0, len(source), 12):
        chunk = source[offset:offset + 12]
        body = ", ".join(f"0x{byte:02x}" for byte in chunk)
        lines.append(f"    {body},")
    lines.append("};")
    lines.append(r"""
int main(int argc, char **argv) {
    const char *volatile keep = tool_marker;
    char path[] = "/tmp/checkdepend-embed-XXXXXX";
    int fd;
    const unsigned char *ptr;
    size_t left;
    char **args;
    int i;

    if (keep[0] == '\0')
        return 126;
    fd = mkstemp(path);
    if (fd < 0) {
        perror("mkstemp");
        return 127;
    }
    ptr = script;
    left = sizeof(script);
    while (left > 0) {
        ssize_t wrote = write(fd, ptr, left);
        if (wrote < 0) {
            perror("write");
            close(fd);
            unlink(path);
            return 127;
        }
        ptr += wrote;
        left -= (size_t)wrote;
    }
    if (close(fd) != 0) {
        perror("close");
        unlink(path);
        return 127;
    }

    /* interpreter, script path, the original argv[1..], and a NULL terminator. */
    args = calloc((size_t)argc + 2, sizeof(char *));
    if (!args) {
        unlink(path);
        return 127;
    }
    args[0] = (char *)interpreter;
    args[1] = path;
    for (i = 1; i < argc; i++)
        args[i + 1] = argv[i];
    execv(interpreter, args);
    perror("execv");
    unlink(path);
    return 127;
}
""")
    return "\n".join(lines) + "\n"

class CompileError(Exception):
    pass

def build_binary() -> bytes:
    if shutil.which("gcc") is None:
        raise CompileError("gcc is required to compile the binary.")
    try:
        source = Path(__file__).read_bytes()
    except OSError as exc:
        raise CompileError(f"Cannot read {__file__}: {exc}") from exc
    if not source:
        raise CompileError(f"{__file__} is empty.")

    with tempfile.TemporaryDirectory(prefix="checkdepend-build-") as tmp:
        c_path = Path(tmp) / "checkdepend.c"
        bin_path = Path(tmp) / INSTALL_NAME
        c_path.write_text(render_c(source, sys.executable), encoding="utf-8")
        proc = subprocess.run(
            ["gcc", "-O2", "-Wall", "-Wextra", "-o", str(bin_path), str(c_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        if proc.returncode != 0:
            detail = proc.stderr.decode("utf-8", errors="replace").strip()
            raise CompileError(detail or "gcc failed")
        return bin_path.read_bytes()

def choose_dest() -> Path:
    system_dir = Path("/usr/local/bin")
    user_dir = Path.home() / ".local" / "bin"
    if os.access(system_dir, os.W_OK):
        return system_dir / INSTALL_NAME
    user_dir.mkdir(parents=True, exist_ok=True)
    if not os.access(user_dir, os.W_OK):
        raise CompileError(f"Cannot write to {system_dir} or {user_dir}.")
    return user_dir / INSTALL_NAME

def candidate_install_paths():
    paths = []
    if STAMP.exists():
        try:
            recorded = STAMP.read_text(encoding="utf-8").strip()
        except OSError:
            recorded = ""
        if recorded:
            paths.append(Path(recorded))
    paths.append(Path("/usr/local/bin") / INSTALL_NAME)
    paths.append(Path.home() / ".local" / "bin" / INSTALL_NAME)
    unique = []
    seen = set()
    for path in paths:
        key = str(path)
        if key not in seen:
            seen.add(key)
            unique.append(path)
    return unique

def is_ours(path: Path) -> bool:
    try:
        if not path.is_file() or path.stat().st_size > 20 * 1024 * 1024:
            return False
        return TOOL_MARKER.encode("ascii") in path.read_bytes()
    except OSError:
        return False

def remove_installed() -> list:
    """Remove installed copies of this tool. Return the paths that were removed."""
    removed = []
    for path in candidate_install_paths():
        if path.exists() and is_ours(path):
            path.unlink()
            removed.append(path)
    if STAMP.exists():
        try:
            STAMP.unlink()
        except OSError:
            pass
    return removed

def place_binary(blob: bytes, dest: Path):
    dest.parent.mkdir(parents=True, exist_ok=True)
    temporary = dest.with_name(dest.name + ".tmp")
    temporary.write_bytes(blob)
    os.chmod(temporary, 0o755)
    os.replace(temporary, dest)
    STAMP.parent.mkdir(parents=True, exist_ok=True)
    STAMP.write_text(str(dest) + "\n", encoding="utf-8")

def install_binary(word: str) -> int:
    try:
        dest = choose_dest()
        if dest.exists() and not is_ours(dest):
            print(f"{dest} exists and is not this tool. Not replacing it.")
            return 1
        print("Compiling binary...")
        blob = build_binary()
        place_binary(blob, dest)
        for path in candidate_install_paths():
            if os.path.normpath(str(path)) == os.path.normpath(str(dest)):
                continue
            if path.exists() and is_ours(path):
                path.unlink()
    except CompileError as exc:
        print(str(exc))
        return 1
    except OSError as exc:
        print(f"Install failed: {exc}")
        return 1

    print(f"{word} {dest}")
    path_entries = os.environ.get("PATH", "").split(":")
    if str(dest.parent) not in path_entries:
        print(f"{dest.parent} is not on PATH. Add it, or run {dest} directly.")
    return 0

def uninstall_binary() -> int:
    try:
        removed = remove_installed()
    except OSError as exc:
        print(f"Uninstall failed: {exc}")
        return 1
    if not removed:
        print("checkdepend is not installed.")
        return 1
    for path in removed:
        print(f"Removed {path}")
    return 0

# ---------- entry ----------
def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        prog=program_name(),
        description="Find upstream Fixes patches, or check backport dependencies.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""\
modes:
  ./checkdepend.py
      ask for a Linux upstream tree, paste commit IDs, update master,
      and list commits that carry a Fixes: tag for each one

  ./checkdepend.py --backport
      check whether fixes for commits in your tree are already applied

  ./checkdepend.py --install
  ./checkdepend.py --uninstall
  ./checkdepend.py --reinstall
      compile a binary with gcc and install or remove it
""",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--backport", action="store_true",
        help="run the dependency check against your tree and a Torvalds tree",
    )
    mode.add_argument(
        "--install", action="store_true",
        help="compile a binary and install it on PATH",
    )
    mode.add_argument(
        "--uninstall", action="store_true",
        help="remove the installed binary",
    )
    mode.add_argument(
        "--reinstall", action="store_true",
        help="compile and install again, replacing the previous binary",
    )
    args = parser.parse_args(argv)

    if args.install:
        return install_binary("Installed")
    if args.uninstall:
        return uninstall_binary()
    if args.reinstall:
        return install_binary("Reinstalled")
    if args.backport:
        ask_excel(run_backport(), "backport")
        return 0

    status, rows = run_fixes()
    if rows:
        ask_excel(rows, "fixes")
    return status


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print()
        sys.exit(130)
